import os
import re
import threading
import ctypes
import tempfile
import webbrowser
from collections import Counter
from datetime import datetime
import tkinter as tk
from tkinter import filedialog, messagebox, scrolledtext, ttk
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Cm, Pt, RGBColor
import pdfplumber

try:
    ctypes.windll.shcore.SetProcessDpiAwareness(2)
except Exception:
    try:
        ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass

CURSOS_PERMITIDOS = [
    "CLASSE ESPECIAL D.I.",
    "ENSINO FUND.1/5 ANO-SERIE",
    "EDUC INFANTIL",
]

def aplicar_cor_fundo_celula(celula, hex_color="E0E0E0"):
    """Aplica cor de fundo (shading) numa célula da tabela do Word."""
    tcPr = celula._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)


def aplicar_no_wrap_celula(celula):
    """Aplica a propriedade XML w:noWrap para proibir estritamente quebra de linha na célula."""
    tcPr = celula._tc.get_or_add_tcPr()
    noWrap_xml = parse_xml(f'<w:noWrap {nsdecls("w")}/>')
    if tcPr.find(noWrap_xml.tag) is None:
        tcPr.append(noWrap_xml)


def formatar_paragrafo_celula(
    p, texto, alignment, bold=False, color=None, font_size=12
):
    """Aplica texto e formatação compacta numa célula."""
    p.text = texto
    p.alignment = alignment
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.0

    for run in p.runs:
        run.font.name = "Arial"
        run.font.size = Pt(font_size)
        run.bold = bold
        if color:
            run.font.color.rgb = color

def formatar_titulo_turma(curso, seriacao, letra_turma):
    """Gera o título do cabeçalho da turma conforme as regras."""
    curso_upper = curso.upper().strip()
    if "CLASSE ESPECIAL D.I." in curso_upper:
        return f"D.I. - Turma {letra_turma}"
    else:
        return f"{seriacao} - Turma {letra_turma}"


def normalizar_turno(turno_raw):
    """Padroniza o nome do turno para Manhã, Tarde, Noite, etc."""
    if not turno_raw:
        return "Não informado"
    t = turno_raw.upper().strip()
    if "MANH" in t or "MATUTIN" in t:
        return "Manhã"
    if "TARD" in t or "VESPERTIN" in t:
        return "Tarde"
    if "NOIT" in t or "NOTURN" in t:
        return "Noite"
    if "INTEGRAL" in t:
        return "Integral"
    return turno_raw.strip().capitalize()


def formatar_nome_serie(seriacao):
    """Extrai uma representação limpa da série (ex: '1º', '2º', 'Infantil 4', 'Classe Especial D.I.')."""
    seriacao_upper = seriacao.upper().strip()
    if any(k in seriacao_upper for k in ["INFANTIL", "PRE", "PRÉ", "MATERNAL"]):
        return seriacao.strip()

    match = re.search(r"(\d+[º°]?)", seriacao)
    if match:
        num_serie = match.group(1)
        if not ("º" in num_serie or "°" in num_serie):
            num_serie += "º"
        return num_serie
    return seriacao.strip()


def formatar_nome_curto_turma(seriacao, letra_turma, turma_titulo):
    """Gera um nome curto para a turma (ex: '1º A', '2º B', 'Infantil 4 A', 'D.I. - Turma A')."""
    seriacao_upper = seriacao.upper().strip()
    if any(k in seriacao_upper for k in ["INFANTIL", "PRE", "PRÉ", "MATERNAL"]):
        return f"{seriacao.strip()} {letra_turma.strip()}".strip()

    match = re.search(r"(\d+[º°]?)", seriacao)
    if match and letra_turma:
        num_serie = match.group(1)
        if not ("º" in num_serie or "°" in num_serie):
            num_serie += "º"
        return f"{num_serie} {letra_turma.strip()}"
    return turma_titulo.strip()


def chave_ordenacao_natural(texto):
    """Chave para ordenação natural de strings com números (ex: 1º, 2º... 10º)."""
    return [int(c) if c.isdigit() else c.lower() for c in re.split(r"(\d+)", str(texto))]


def ordenar_turnos(turno_nome):
    """Ordena turnos dando prioridade a Manhã, depois Tarde e demais."""
    t = str(turno_nome).upper().strip()
    if "MANH" in t or "MATUTIN" in t:
        return (0, t)
    if "TARD" in t or "VESPERTIN" in t:
        return (1, t)
    if "INTEGRAL" in t:
        return (2, t)
    if "NOIT" in t or "NOTURN" in t:
        return (3, t)
    return (4, t)

def estimar_largura_texto_cm(texto, font_size_pt=12):
    """Calcula uma estimativa da largura em cm de uma string no estilo Arial."""
    if not texto:
        return 0.0

    total_pts = 0.0
    for char in str(texto):
        if char in "WMwmQ":
            total_pts += 0.85
        elif char.isupper():
            if char in "IJLT":
                total_pts += 0.42
            elif char in "FEE":
                total_pts += 0.60
            else:
                total_pts += 0.72
        elif char.islower():
            if char in "ijl":
                total_pts += 0.28
            elif char in "ftr":
                total_pts += 0.38
            else:
                total_pts += 0.55
        elif char.isdigit():
            total_pts += 0.55
        elif char in " -_./()":
            total_pts += 0.35
        else:
            total_pts += 0.50

    largura_cm = (total_pts * font_size_pt * 2.54) / 72.0
    return largura_cm


def detectar_estilo_case(turmas_dados):
    """Detecta se os nomes nas turmas do Word estão em Title Case ('title') ou Maiúsculas ('upper')."""
    total_nomes = 0
    com_minusculas = 0

    for alunos in turmas_dados.values():
        for info in alunos.values():
            nome = info.get("nome", "").strip()
            if nome:
                total_nomes += 1
                if any(c.islower() for c in nome):
                    com_minusculas += 1

    if total_nomes > 0 and (com_minusculas / total_nomes) > 0.3:
        return "title"
    return "upper"


def calcular_larguras_colunas_turma(
    alunos, colunas_extras, opcao_case="upper", largura_total_cm=18.0, font_size_nome=12
):
    """
    Calcula as larguras das colunas para uma turma específica.
    Quando há colunas extras, a largura de 'Nome do Estudante' é dimensionada
    exatamente ao maior nome daquela turma (+ margem de segurança), sem quebras de linha.
    """
    w_num = 1.0   # Largura N.º
    qtd_extras = len(colunas_extras)

    if qtd_extras > 0:
        max_largura_nome = estimar_largura_texto_cm("Nome do Estudante", font_size_pt=12) + 0.7

        for num, aluno in alunos.items():
            nome = aluno.get("nome", "")
            if opcao_case == "upper":
                nome = nome.upper()
            elif opcao_case == "title":
                if any(c.islower() for c in nome):
                    pass
                else:
                    nome = nome.title()

            largura_nome = estimar_largura_texto_cm(nome, font_size_pt=font_size_nome) + 0.7
            if largura_nome > max_largura_nome:
                max_largura_nome = largura_nome

        w_nome = max_largura_nome
        espaco_restante = largura_total_cm - (w_num + w_nome)

        min_w_extra = 1.2
        if espaco_restante < (qtd_extras * min_w_extra):
            w_nome = max(4.0, largura_total_cm - w_num - (qtd_extras * min_w_extra))
            espaco_restante = largura_total_cm - (w_num + w_nome)

        w_col_extra = espaco_restante / qtd_extras
        w_extras = [w_col_extra] * qtd_extras
    else:
        w_nome = largura_total_cm - w_num
        w_extras = []

    larguras_finais = [Cm(w_num), Cm(w_nome)] + [Cm(w) for w in w_extras]
    return larguras_finais

def extrair_dados_pdf(caminho_pdf, log_callback=None, progress_callback=None):
    """Extrai turmas e alunos do PDF filtrando os cursos permitidos."""
    turmas = {}
    turma_atual = None
    processar_turma_atual = False

    curso_atual = ""
    seriacao_atual = ""
    turno_atual = ""
    letra_atual = ""

    nome_arquivo = os.path.basename(caminho_pdf)
    if log_callback:
        log_callback(f"Lendo {nome_arquivo}")

    with pdfplumber.open(caminho_pdf) as pdf:
        total_paginas = len(pdf.pages)
        for idx, pagina in enumerate(pdf.pages):
            if progress_callback and total_paginas > 0:
                progress_callback((idx + 1) / total_paginas)

            texto = pagina.extract_text()
            if not texto:
                continue

            linhas = texto.split("\n")
            for linha in linhas:
                linha_str = linha.strip()

                if (
                    "Curso:" in linha_str
                    or "Seriação:" in linha_str
                    or "Turma:" in linha_str
                ):
                    match_turma = re.search(
                        r"Curso:\s*(.*?)\s+Seriação:\s*(.*?)(?:\s+Turno:\s*(.*?))?\s+Turma:\s*(.*)",
                        linha_str,
                        re.IGNORECASE,
                    )
                    if match_turma:
                        curso = match_turma.group(1).strip()
                        seriacao = match_turma.group(2).strip()
                        turno_raw = match_turma.group(3).strip() if match_turma.group(3) else ""
                        letra_turma = match_turma.group(4).strip()

                        curso_permitido = any(
                            c in curso.upper() for c in CURSOS_PERMITIDOS
                        )

                        if curso_permitido:
                            processar_turma_atual = True
                            turma_atual = formatar_titulo_turma(
                                curso, seriacao, letra_turma
                            )
                            curso_atual = curso
                            seriacao_atual = seriacao
                            turno_atual = normalizar_turno(turno_raw)
                            letra_atual = letra_turma

                            if turma_atual not in turmas:
                                turmas[turma_atual] = {}
                        else:
                            processar_turma_atual = False
                            turma_atual = None
                    continue

                if processar_turma_atual and turma_atual:
                    match_aluno = re.match(
                        r"^\s*(\d+)\s+\d+\s+(.+?)\s+\d{2}/\d{2}/\d{4}.*?\b(Matriculado|Transferido|Remanejado)\b",
                        linha_str,
                        re.IGNORECASE,
                    )

                    if match_aluno:
                        num_chamada = int(match_aluno.group(1))
                        nome_aluno = match_aluno.group(2).strip()
                        situacao = match_aluno.group(3).strip().capitalize()

                        turmas[turma_atual][num_chamada] = {
                            "nome": nome_aluno,
                            "situacao": situacao,
                            "turno": turno_atual,
                            "seriacao": seriacao_atual,
                            "letra_turma": letra_atual,
                            "curso": curso_atual,
                        }

    return turmas

def calcular_totais_dados(dados_consolidados):
    """Calcula os totais e agrupa turmas por turno (Manhã, Tarde, etc.)."""
    total_matriculados = 0
    por_curso = {}
    por_turno = {}
    por_serie = {}
    por_turma = {}
    por_turma_agrupado = {}

    for turma_titulo, alunos in dados_consolidados.items():
        count_turma_matriculados = 0
        sample_aluno = None

        for num, aluno in alunos.items():
            situacao = aluno.get("situacao", "").strip().capitalize()
            if situacao == "Matriculado":
                if sample_aluno is None:
                    sample_aluno = aluno

                total_matriculados += 1
                count_turma_matriculados += 1

                curso = aluno.get("curso", "Não informado")
                turno = aluno.get("turno", "Não informado")
                seriacao = aluno.get("seriacao", "Outros")

                nome_serie = formatar_nome_serie(seriacao)

                por_curso[curso] = por_curso.get(curso, 0) + 1
                por_turno[turno] = por_turno.get(turno, 0) + 1
                por_serie[nome_serie] = por_serie.get(nome_serie, 0) + 1

        if not sample_aluno and alunos:
            sample_aluno = next(iter(alunos.values()))

        if sample_aluno:
            seriacao = sample_aluno.get("seriacao", "")
            letra_turma = sample_aluno.get("letra_turma", "")
            turno_turma = sample_aluno.get("turno", "Não informado")
            if not turno_turma:
                turno_turma = "Não informado"

            nome_curto_turma = formatar_nome_curto_turma(
                seriacao, letra_turma, turma_titulo
            )

            if turno_turma and turno_turma != "Não informado":
                rotulo_turma = f"{nome_curto_turma} ({turno_turma})"
            else:
                rotulo_turma = nome_curto_turma

            por_turma[rotulo_turma] = count_turma_matriculados

            if turno_turma not in por_turma_agrupado:
                por_turma_agrupado[turno_turma] = {}
            por_turma_agrupado[turno_turma][nome_curto_turma] = count_turma_matriculados

    return {
        "total_geral": total_matriculados,
        "por_curso": por_curso,
        "por_turno": por_turno,
        "por_serie": por_serie,
        "por_turma": por_turma,
        "por_turma_agrupado": por_turma_agrupado,
    }

def ler_dados_word_existente(caminho_word):
    """Lê um ficheiro Word existente, extrai os dados, colunas extras e o tamanho da fonte usada nos nomes."""
    if not os.path.exists(caminho_word):
        return {}, [], 12

    doc = Document(caminho_word)
    turmas_existentes = {}
    colunas_extras = []
    tamanhos_fonte_encontrados = []

    if doc.tables:
        primeira_tabela = doc.tables[0]
        num_cols = len(primeira_tabela.rows[0].cells)
        if num_cols > 2:
            for c_idx in range(2, num_cols):
                colunas_extras.append(primeira_tabela.rows[0].cells[c_idx].text.strip())

    tabela_idx = 0
    for p in doc.paragraphs:
        texto = p.text.strip()
        if texto:
            turma_atual = texto
            if turma_atual not in turmas_existentes:
                turmas_existentes[turma_atual] = {}

            if tabela_idx < len(doc.tables):
                tabela = doc.tables[tabela_idx]
                for row in tabela.rows[1:]:
                    if not row.cells:
                        continue
                    n_text = row.cells[0].text.strip()
                    nome_text = row.cells[1].text.strip()

                    # Inspeciona a fonte do texto do nome do aluno
                    if len(row.cells) > 1 and row.cells[1].paragraphs:
                        for run in row.cells[1].paragraphs[0].runs:
                            if run.font and run.font.size is not None:
                                tamanhos_fonte_encontrados.append(run.font.size.pt)

                    situacao = "Matriculado"

                    if len(row.cells) >= 3:
                        for cell_extra in row.cells[2:]:
                            ass_text = cell_extra.text.strip().lower()
                            if "transferido" in ass_text:
                                situacao = "Transferido"
                                break
                            elif "remanejado" in ass_text:
                                situacao = "Remanejado"
                                break

                    if "transferido" in nome_text.lower():
                        situacao = "Transferido"
                        nome_text = re.sub(
                            r"\s*-\s*transferido",
                            "",
                            nome_text,
                            flags=re.IGNORECASE,
                        ).strip()
                    elif "remanejado" in nome_text.lower():
                        situacao = "Remanejado"
                        nome_text = re.sub(
                            r"\s*-\s*remanejado",
                            "",
                            nome_text,
                            flags=re.IGNORECASE,
                        ).strip()

                    if n_text.isdigit():
                        num = int(n_text)
                        turmas_existentes[turma_atual][num] = {
                            "nome": nome_text,
                            "situacao": situacao,
                        }
                tabela_idx += 1

    fonte_detectada = 12
    if tamanhos_fonte_encontrados:
        fonte_mais_comum = Counter(tamanhos_fonte_encontrados).most_common(1)[0][0]
        fonte_detectada = int(round(fonte_mais_comum))

    return turmas_existentes, colunas_extras, fonte_detectada

def gerar_documento_word(
    turmas_dados,
    caminho_saida,
    colunas_extras=None,
    opcao_case="upper",
    font_size_nome=12,
    log_callback=None,
    progress_callback=None,
):
    """Gera o documento Word ajustando dinamicamente a largura dos nomes de cada turma sem quebrar linha."""
    if colunas_extras is None:
        colunas_extras = []

    doc = Document()

    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.5)
        section.right_margin = Cm(1.5)

    primeira_turma = True
    num_cols = 2 + len(colunas_extras)
    titulos = ["N.º", "Nome do Estudante"] + colunas_extras

    if log_callback:
        log_callback("Gerando tabelas no Word...")

    turmas_chaves = sorted(turmas_dados.keys())
    total_turmas = len(turmas_chaves)

    for idx_turma, turma_titulo in enumerate(turmas_chaves):
        if progress_callback and total_turmas > 0:
            progress_callback((idx_turma + 1) / total_turmas)

        alunos = turmas_dados[turma_titulo]
        if not alunos:
            continue

        if not primeira_turma:
            doc.add_page_break()
        primeira_turma = False

        larguras = calcular_larguras_colunas_turma(
            alunos,
            colunas_extras,
            opcao_case=opcao_case,
            largura_total_cm=18.0,
            font_size_nome=font_size_nome,
        )

        p_titulo = doc.add_paragraph()
        p_titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_titulo.paragraph_format.space_after = Pt(6)
        run_t = p_titulo.add_run(turma_titulo)
        run_t.bold = True
        run_t.font.size = Pt(14)
        run_t.font.name = "Arial"

        tabela = doc.add_table(rows=1, cols=num_cols)
        tabela.style = "Table Grid"
        tabela.alignment = WD_TABLE_ALIGNMENT.CENTER
        tabela.autofit = False

        for i, col in enumerate(tabela.columns):
            col.width = larguras[i]

        hdr_cells = tabela.rows[0].cells
        for i, titulo in enumerate(titulos):
            hdr_cells[i].width = larguras[i]
            aplicar_cor_fundo_celula(hdr_cells[i], "2F4F4F")
            aplicar_no_wrap_celula(hdr_cells[i])

            p = hdr_cells[i].paragraphs[0]
            align = (
                WD_ALIGN_PARAGRAPH.CENTER
                if i != 1
                else WD_ALIGN_PARAGRAPH.LEFT
            )
            formatar_paragrafo_celula(
                p,
                titulo,
                align,
                bold=True,
                color=RGBColor(255, 255, 255),
                font_size=12,
            )

        for num in sorted(alunos.keys()):
            aluno = alunos[num]
            row_cells = tabela.add_row().cells

            aplicar_no_wrap_celula(row_cells[1])

            for i in range(num_cols):
                row_cells[i].width = larguras[i]
                row_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

            situacao_str = aluno["situacao"].capitalize()
            is_especial = situacao_str in ["Transferido", "Remanejado"]

            nome_formatado = aluno["nome"]
            if opcao_case == "upper":
                nome_formatado = nome_formatado.upper()
            elif opcao_case == "title":
                if any(c.islower() for c in nome_formatado):
                    pass
                else:
                    nome_formatado = nome_formatado.title()

            formatar_paragrafo_celula(
                row_cells[0].paragraphs[0],
                str(num),
                WD_ALIGN_PARAGRAPH.CENTER,
                bold=True,
                font_size=12,
            )

            if is_especial:
                if len(colunas_extras) >= 2:
                    formatar_paragrafo_celula(
                        row_cells[1].paragraphs[0],
                        nome_formatado,
                        WD_ALIGN_PARAGRAPH.LEFT,
                        bold=False,
                        font_size=font_size_nome,
                    )
                    celula_mesclada_extras = row_cells[2].merge(row_cells[-1])
                    celula_mesclada_extras.width = sum(larguras[2:])
                    formatar_paragrafo_celula(
                        celula_mesclada_extras.paragraphs[0],
                        situacao_str,
                        WD_ALIGN_PARAGRAPH.CENTER,
                        bold=True,
                        font_size=12,
                    )
                elif len(colunas_extras) == 1:
                    formatar_paragrafo_celula(
                        row_cells[1].paragraphs[0],
                        nome_formatado,
                        WD_ALIGN_PARAGRAPH.LEFT,
                        bold=False,
                        font_size=font_size_nome,
                    )
                    formatar_paragrafo_celula(
                        row_cells[2].paragraphs[0],
                        situacao_str,
                        WD_ALIGN_PARAGRAPH.CENTER,
                        bold=True,
                        font_size=12,
                    )
                else:
                    p_nome = row_cells[1].paragraphs[0]
                    p_nome.text = ""
                    p_nome.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    p_nome.paragraph_format.space_before = Pt(2)
                    p_nome.paragraph_format.space_after = Pt(2)
                    p_nome.paragraph_format.line_spacing = 1.0

                    run_nome = p_nome.add_run(nome_formatado)
                    run_nome.font.name = "Arial"
                    run_nome.font.size = Pt(font_size_nome)
                    run_nome.bold = False

                    run_status = p_nome.add_run(f" - {situacao_str}")
                    run_status.font.name = "Arial"
                    run_status.font.size = Pt(12)
                    run_status.bold = True

                for cell in row_cells:
                    aplicar_cor_fundo_celula(cell, "E0E0E0")
            else:
                formatar_paragrafo_celula(
                    row_cells[1].paragraphs[0],
                    nome_formatado,
                    WD_ALIGN_PARAGRAPH.LEFT,
                    bold=False,
                    font_size=font_size_nome,
                )

                for col_idx in range(len(colunas_extras)):
                    cell_extra = row_cells[2 + col_idx]
                    formatar_paragrafo_celula(
                        cell_extra.paragraphs[0],
                        "",
                        WD_ALIGN_PARAGRAPH.CENTER,
                        bold=False,
                        font_size=12,
                    )

    doc.save(caminho_saida)
    if log_callback:
        log_callback("Documento salvo com sucesso!")

class SideTabWizardApp:

    def __init__(self, root):
        self.root = root
        self.root.title("Gerador de Listas de Alunos")
        self.root.protocol("WM_DELETE_WINDOW", self._ao_fechar)

        self._centralizar_janela(1150, 750)
        self.root.minsize(1000, 650)

        # Dados da Aplicação
        self.opcao_modo = tk.StringVar(value="1")
        self.opcao_case = tk.StringVar(value="upper")
        self.fonte_nome_11 = tk.BooleanVar(value=False)
        self.fonte_detectada_word = 12  # Armazena o tamanho da fonte ao atualizar Word existente
        self.caminho_word = ""
        self.colunas_extras = []
        self.caminhos_pdfs = []
        self.arquivos_temporarios = []

        self.passo_atual = 1
        self.max_passo_alcancado = 1

        self.passos_info = [
            (1, "Modo de Operação"),
            (2, "Arquivo Word"),
            (3, "Layout da Tabela"),
            (4, "Seleção de PDFs"),
            (5, "Execução"),
        ]

        self._configurar_estilos()
        self._construir_layout_principal()
        self._criar_passos_conteudo()
        self._exibir_passo(1)
        self._atualizar_resumo()

    def _ao_fechar(self):
        """Remove arquivos temporários criados ao fechar a aplicação."""
        for caminho in self.arquivos_temporarios:
            try:
                if os.path.exists(caminho):
                    os.remove(caminho)
            except Exception:
                pass
        self.root.destroy()

    def _centralizar_janela(self, largura, altura):
        self.root.update_idletasks()
        largura_tela = self.root.winfo_screenwidth()
        altura_tela = self.root.winfo_screenheight()

        pos_x = max(0, (largura_tela // 2) - (largura // 2))
        pos_y = max(0, (altura_tela // 2) - (altura // 2))

        self.root.geometry(f"{largura}x{altura}+{pos_x}+{pos_y}")

    def _configurar_estilos(self):
        self.style = ttk.Style()

        self.style.configure("SidebarTitle.TLabel", font=("TkDefaultFont", 10, "bold"))
        self.style.configure("Header.TLabel", font=("TkDefaultFont", 12, "bold"))
        self.style.configure("SubHeader.TLabel", font=("TkDefaultFont", 9))
        self.style.configure("SummaryHeader.TLabel", font=("TkDefaultFont", 10, "bold"))
        self.style.configure("SummaryLabel.TLabel", font=("TkDefaultFont", 9, "bold"))
        self.style.configure("SummaryVal.TLabel", font=("TkDefaultFont", 9), wraplength=200)

        self.style.configure("StepItem.TLabel", font=("TkDefaultFont", 9), padding=(8, 6))
        self.style.configure("StepItemActive.TLabel", font=("TkDefaultFont", 9, "bold"), padding=(8, 6))

    def _construir_layout_principal(self):
        self.container = ttk.Frame(self.root)
        self.container.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)

        # 1. SIDEBAR ESQUERDA
        self.sidebar_left = ttk.Frame(self.container, width=230)
        self.sidebar_left.pack(side=tk.LEFT, fill=tk.Y, expand=False, padx=(5, 5))
        self.sidebar_left.pack_propagate(False)

        lbl_sidebar_title = ttk.Label(
            self.sidebar_left, text="Passos do Processo", style="SidebarTitle.TLabel"
        )
        lbl_sidebar_title.pack(anchor=tk.W, padx=10, pady=(15, 10))

        ttk.Separator(self.sidebar_left, orient="horizontal").pack(fill=tk.X, padx=5, pady=(0, 10))

        self.labels_sidebar = {}
        for num, nome in self.passos_info:
            lbl = ttk.Label(
                self.sidebar_left, text=f"  {num}. {nome}", style="StepItem.TLabel", cursor="hand2"
            )
            lbl.pack(fill=tk.X, padx=5, pady=2)
            lbl.bind("<Button-1>", lambda e, p=num: self._clique_aba_lateral(p))
            self.labels_sidebar[num] = lbl

        btn_reiniciar = ttk.Button(self.sidebar_left, text="⟲ Reiniciar Tudo", command=self.reiniciar)
        btn_reiniciar.pack(side=tk.BOTTOM, fill=tk.X, padx=5, pady=10)

        ttk.Separator(self.container, orient="vertical").pack(side=tk.LEFT, fill=tk.Y, padx=5)

        # 2. PAINEL DE RESUMO À DIREITA
        self.summary_right = ttk.Frame(self.container, width=250)
        self.summary_right.pack(side=tk.RIGHT, fill=tk.Y, expand=False, padx=(5, 5))
        self.summary_right.pack_propagate(False)

        ttk.Separator(self.container, orient="vertical").pack(side=tk.RIGHT, fill=tk.Y, padx=5)

        lbl_summary_title = ttk.Label(
            self.summary_right, text="Resumo das Escolhas", style="SummaryHeader.TLabel"
        )
        lbl_summary_title.pack(anchor=tk.W, padx=10, pady=(15, 10))

        ttk.Separator(self.summary_right, orient="horizontal").pack(fill=tk.X, padx=5, pady=(0, 10))

        ttk.Label(self.summary_right, text="Modo:", style="SummaryLabel.TLabel").pack(anchor=tk.W, padx=10)
        self.sum_val_modo = ttk.Label(self.summary_right, text="-", style="SummaryVal.TLabel")
        self.sum_val_modo.pack(anchor=tk.W, padx=10, pady=(0, 10))

        ttk.Label(self.summary_right, text="Arquivo Word:", style="SummaryLabel.TLabel").pack(anchor=tk.W, padx=10)
        self.sum_val_word = ttk.Label(self.summary_right, text="Não selecionado", style="SummaryVal.TLabel")
        self.sum_val_word.pack(anchor=tk.W, padx=10, pady=(0, 10))

        ttk.Label(self.summary_right, text="Estilo dos Nomes:", style="SummaryLabel.TLabel").pack(anchor=tk.W, padx=10)
        self.sum_val_case = ttk.Label(self.summary_right, text="TUDO MAIÚSCULO (AA)", style="SummaryVal.TLabel")
        self.sum_val_case.pack(anchor=tk.W, padx=10, pady=(0, 10))

        ttk.Label(self.summary_right, text="Layout da Tabela:", style="SummaryLabel.TLabel").pack(anchor=tk.W, padx=10)
        self.sum_val_layout = ttk.Label(self.summary_right, text="-", style="SummaryVal.TLabel")
        self.sum_val_layout.pack(anchor=tk.W, padx=10, pady=(0, 10))

        ttk.Label(self.summary_right, text="Arquivos PDF:", style="SummaryLabel.TLabel").pack(anchor=tk.W, padx=10)
        self.sum_val_pdfs = ttk.Label(self.summary_right, text="Nenhum PDF selecionado", style="SummaryVal.TLabel")
        self.sum_val_pdfs.pack(anchor=tk.W, padx=10, pady=(0, 10))

        # 3. ÁREA DE CONTEÚDO CENTRAL
        self.center_content = ttk.Frame(self.container, padding="15")
        self.center_content.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)

        self.frames_passos = {}

    def _criar_passos_conteudo(self):
        # Passo 1: Modo
        f1 = ttk.Frame(self.center_content)
        ttk.Label(f1, text="Passo 1: Modo de Operação", style="Header.TLabel").pack(anchor=tk.W)
        ttk.Label(
            f1,
            text="Escolha a operação desejada: criar/atualizar um arquivo Word ou listar os totais das turmas.",
            style="SubHeader.TLabel",
        ).pack(anchor=tk.W, pady=(2, 15))

        f1_card = ttk.LabelFrame(f1, text=" Opções ", padding="15")
        f1_card.pack(fill=tk.X, pady=10)

        ttk.Radiobutton(
            f1_card,
            text="Criar um NOVO arquivo Word",
            variable=self.opcao_modo,
            value="1",
            command=self._atualizar_resumo,
        ).pack(anchor=tk.W, pady=5)
        ttk.Radiobutton(
            f1_card,
            text="ATUALIZAR um arquivo Word existente",
            variable=self.opcao_modo,
            value="2",
            command=self._atualizar_resumo,
        ).pack(anchor=tk.W, pady=5)
        ttk.Radiobutton(
            f1_card,
            text="Listar os números totais das turmas",
            variable=self.opcao_modo,
            value="3",
            command=self._atualizar_resumo,
        ).pack(anchor=tk.W, pady=5)

        self._criar_bar_navegacao(f1, btn_proximo_cmd=lambda: self._avancar_passo(1))
        self.frames_passos[1] = f1

        # Passo 2: Arquivo Word
        f2 = ttk.Frame(self.center_content)
        self.lbl_p2_titulo = ttk.Label(f2, text="Passo 2: Arquivo Word", style="Header.TLabel")
        self.lbl_p2_titulo.pack(anchor=tk.W)
        self.lbl_p2_desc = ttk.Label(f2, text="Selecione o destino do arquivo Word.", style="SubHeader.TLabel")
        self.lbl_p2_desc.pack(anchor=tk.W, pady=(2, 15))

        f2_card = ttk.LabelFrame(f2, text=" Seleção de Arquivo ", padding="15")
        f2_card.pack(fill=tk.X, pady=10)

        self.lbl_word_selected = ttk.Label(f2_card, text="Nenhum arquivo selecionado.", wraplength=550)
        self.lbl_word_selected.pack(side=tk.LEFT, fill=tk.X, expand=True)

        ttk.Button(f2_card, text="Abrir...", command=self._selecionar_word).pack(side=tk.RIGHT)

        self._criar_bar_navegacao(
            f2,
            btn_voltar_cmd=lambda: self._voltar_passo(2),
            btn_proximo_cmd=lambda: self._avancar_passo(2),
        )
        self.frames_passos[2] = f2

        # Passo 3: Layout Tabela e Formatação do Nome
        f3 = ttk.Frame(self.center_content)
        ttk.Label(f3, text="Passo 3: Layout da Tabela e Nomes", style="Header.TLabel").pack(anchor=tk.W)
        ttk.Label(
            f3,
            text="Configure a formatação dos nomes e adicione/remova colunas da sua tabela abaixo:",
            style="SubHeader.TLabel",
        ).pack(anchor=tk.W, pady=(2, 15))

        f3_case_card = ttk.LabelFrame(f3, text=" Formatação dos Nomes dos Estudantes ", padding="12")
        f3_case_card.pack(fill=tk.X, pady=(0, 10))

        ttk.Label(
            f3_case_card,
            text="Escolha o estilo de texto para os nomes dos estudantes:",
        ).pack(anchor=tk.W, pady=(0, 6))

        frame_btns_case = ttk.Frame(f3_case_card)
        frame_btns_case.pack(anchor=tk.W)

        fonte_times_bold = ("Times New Roman", 12, "bold italic")

        self.btn_case_upper = tk.Button(
            frame_btns_case,
            text="AA",
            font=fonte_times_bold,
            width=4,
            command=lambda: self._definir_opcao_case("upper"),
        )
        self.btn_case_upper.pack(side=tk.LEFT, padx=(0, 10))

        self.btn_case_title = tk.Button(
            frame_btns_case,
            text="Aa",
            font=fonte_times_bold,
            width=4,
            command=lambda: self._definir_opcao_case("title"),
        )
        self.btn_case_title.pack(side=tk.LEFT)

        self.chk_fonte_11 = ttk.Checkbutton(
            f3_case_card,
            text="Diminuir tamanho da fonte",
            variable=self.fonte_nome_11,
            command=self._atualizar_resumo,
        )
        self.chk_fonte_11.pack(anchor=tk.W, pady=(10, 0))

        self._atualizar_botoes_case()

        f3_card = ttk.LabelFrame(f3, text=" Adicionar Nova Coluna ", padding="12")
        f3_card.pack(fill=tk.X, pady=(0, 10))

        frame_input_col = ttk.Frame(f3_card)
        frame_input_col.pack(fill=tk.X)

        ttk.Label(frame_input_col, text="Título da Coluna:").pack(side=tk.LEFT, padx=(0, 5))
        self.entry_nova_coluna = ttk.Entry(frame_input_col)
        self.entry_nova_coluna.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 5))

        ttk.Button(
            frame_input_col,
            text="➕ Adicionar Coluna",
            command=self._adicionar_coluna_extra,
        ).pack(side=tk.RIGHT)

        f3_list_card = ttk.LabelFrame(f3, text=" Colunas Adicionais Configuradas ", padding="12")
        f3_list_card.pack(fill=tk.BOTH, expand=True, pady=5)

        self.listbox_colunas = tk.Listbox(
            f3_list_card, height=4, font=("TkDefaultFont", 9), selectmode=tk.SINGLE
        )
        self.listbox_colunas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 10))

        btn_box_cols = ttk.Frame(f3_list_card)
        btn_box_cols.pack(side=tk.RIGHT, fill=tk.Y)

        ttk.Button(
            btn_box_cols,
            text="❌ Remover Coluna",
            command=self._remover_coluna_extra,
        ).pack(anchor=tk.N)

        self._criar_bar_navegacao(
            f3,
            btn_voltar_cmd=lambda: self._voltar_passo(3),
            btn_proximo_cmd=lambda: self._avancar_passo(3),
        )
        self.frames_passos[3] = f3

        # Passo 4: Seleção PDFs
        f4 = ttk.Frame(self.center_content)
        ttk.Label(f4, text="Passo 4: Seleção de PDFs", style="Header.TLabel").pack(anchor=tk.W)
        ttk.Label(
            f4,
            text="Escolha um ou mais arquivos PDF com os dados das turmas.",
            style="SubHeader.TLabel",
        ).pack(anchor=tk.W, pady=(2, 15))

        f4_card = ttk.LabelFrame(f4, text=" Arquivos Selecionados ", padding="15")
        f4_card.pack(fill=tk.BOTH, expand=True, pady=10)

        btn_box = ttk.Frame(f4_card)
        btn_box.pack(fill=tk.X, pady=(0, 8))

        ttk.Button(btn_box, text="➕ Adicionar PDFs...", command=self._selecionar_pdfs).pack(
            side=tk.LEFT, padx=(0, 5)
        )
        ttk.Button(
            btn_box,
            text="❌ Remover Selecionado",
            command=self._remover_pdf_selecionado,
        ).pack(side=tk.LEFT)

        self.listbox_pdfs = tk.Listbox(
            f4_card, height=8, font=("TkDefaultFont", 9), selectmode=tk.SINGLE
        )
        self.listbox_pdfs.pack(fill=tk.BOTH, expand=True)

        self._criar_bar_navegacao(
            f4,
            btn_voltar_cmd=lambda: self._voltar_passo(4),
            btn_proximo_cmd=lambda: self._avancar_passo(4),
        )
        self.frames_passos[4] = f4

        # Passo 5: Execução
        f5 = ttk.Frame(self.center_content)
        ttk.Label(f5, text="Passo 5: Processamento e Execução", style="Header.TLabel").pack(anchor=tk.W)
        ttk.Label(
            f5,
            text="Clique no botão para iniciar o processamento.",
            style="SubHeader.TLabel",
        ).pack(anchor=tk.W, pady=(2, 15))

        self.btn_gerar = ttk.Button(
            f5,
            text="⚡ Gerar Documento",
            command=self._executar_processamento,
        )
        self.btn_gerar.pack(fill=tk.X, ipady=5, pady=(0, 10))

        frame_prog = ttk.Frame(f5)
        frame_prog.pack(fill=tk.X, pady=5)

        self.progress = ttk.Progressbar(frame_prog, mode="determinate", maximum=100, value=0)
        self.progress.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=(0, 8))

        self.lbl_porcentagem = ttk.Label(
            frame_prog, text="0%", font=("TkDefaultFont", 9, "bold"), width=5, anchor=tk.E
        )
        self.lbl_porcentagem.pack(side=tk.RIGHT)

        ttk.Label(f5, text="Execução:", font=("TkDefaultFont", 9, "bold")).pack(anchor=tk.W, pady=(10, 2))
        self.txt_log = scrolledtext.ScrolledText(f5, height=12, state="disabled", font=("Consolas", 9))
        self.txt_log.pack(fill=tk.BOTH, expand=True)

        self._criar_bar_navegacao(
            f5,
            btn_voltar_cmd=lambda: self._voltar_passo(5),
            btn_proximo_txt=None,
        )
        self.frames_passos[5] = f5

    def _criar_bar_navegacao(
        self,
        parent_frame,
        btn_voltar_cmd=None,
        btn_proximo_cmd=None,
        btn_proximo_txt="Avançar ➔",
    ):
        nav_frame = ttk.Frame(parent_frame)
        nav_frame.pack(side=tk.BOTTOM, fill=tk.X, pady=(15, 0))

        if btn_voltar_cmd:
            ttk.Button(nav_frame, text="Voltar", command=btn_voltar_cmd).pack(side=tk.LEFT)

        if btn_proximo_cmd and btn_proximo_txt:
            ttk.Button(nav_frame, text=btn_proximo_txt, command=btn_proximo_cmd).pack(side=tk.RIGHT)

    def _definir_opcao_case(self, modo):
        self.opcao_case.set(modo)
        self._atualizar_botoes_case()
        self._atualizar_resumo()

    def _atualizar_botoes_case(self):
        if self.opcao_case.get() == "upper":
            self.btn_case_upper.config(bg="#c5c5c5", fg="black", relief=tk.SUNKEN)
            self.btn_case_title.config(bg="#f0f0f0", fg="black", relief=tk.RAISED)
        else:
            self.btn_case_upper.config(bg="#f0f0f0", fg="black", relief=tk.RAISED)
            self.btn_case_title.config(bg="#c5c5c5", fg="black", relief=tk.SUNKEN)

    def _adicionar_coluna_extra(self):
        titulo = self.entry_nova_coluna.get().strip()
        if not titulo:
            messagebox.showwarning("Atenção", "Digite o título da coluna antes de adicionar.")
            return
        if titulo in self.colunas_extras:
            messagebox.showwarning("Atenção", "Esta coluna já foi adicionada.")
            return

        self.colunas_extras.append(titulo)
        self.entry_nova_coluna.delete(0, tk.END)
        self._atualizar_listbox_colunas()
        self._atualizar_resumo()

    def _remover_coluna_extra(self):
        selecao = self.listbox_colunas.curselection()
        if not selecao:
            messagebox.showwarning("Atenção", "Selecione uma coluna na lista para remover.")
            return

        index = selecao[0]
        del self.colunas_extras[index]
        self._atualizar_listbox_colunas()
        self._atualizar_resumo()

    def _atualizar_listbox_colunas(self):
        self.listbox_colunas.delete(0, tk.END)
        for col in self.colunas_extras:
            self.listbox_colunas.insert(tk.END, col)

    def _exibir_passo(self, passo_num):
        modo = self.opcao_modo.get()

        if modo == "3":
            if passo_num in [2, 3]:
                if self.passo_atual < 2:
                    passo_num = 4
                else:
                    passo_num = 1
        elif modo == "2":
            if passo_num == 3:
                if self.passo_atual < 3:
                    passo_num = 4
                else:
                    passo_num = 2

        self.passo_atual = passo_num
        if passo_num > self.max_passo_alcancado:
            self.max_passo_alcancado = passo_num

        for f in self.frames_passos.values():
            f.pack_forget()

        if passo_num == 2:
            if modo == "1":
                self.lbl_p2_titulo.config(text="Passo 2: Onde salvar o NOVO arquivo Word?")
                self.lbl_p2_desc.config(text="Escolha o local e nome para salvar o documento gerado.")
            elif modo == "2":
                self.lbl_p2_titulo.config(text="Passo 2: Seleção do arquivo Word EXISTENTE")
                self.lbl_p2_desc.config(
                    text="Escolha o arquivo .docx que deseja atualizar com novos dados."
                )

        if passo_num == 5:
            if modo == "3":
                self.btn_gerar.config(text="📊 Calcular Totais das Turmas")
            else:
                self.btn_gerar.config(text="⚡ Gerar Documento Word")

        self.frames_passos[passo_num].pack(fill=tk.BOTH, expand=True)
        self._atualizar_sidebar_visual()

    def _atualizar_sidebar_visual(self):
        modo = self.opcao_modo.get()

        for num, lbl in self.labels_sidebar.items():
            nome_passo = self.passos_info[num - 1][1]

            if modo == "3" and num in [2, 3]:
                lbl.config(style="StepItem.TLabel", text=f"  {num}. {nome_passo} (Inativo)")
                continue

            if modo == "2" and num == 3:
                lbl.config(style="StepItem.TLabel", text=f"  {num}. {nome_passo} (Inativo)")
                continue

            if num == self.passo_atual:
                lbl.config(style="StepItemActive.TLabel", text=f"▶ {num}. {nome_passo}")
            elif num <= self.max_passo_alcancado:
                lbl.config(style="StepItem.TLabel", text=f"✓ {num}. {nome_passo}")
            else:
                lbl.config(style="StepItem.TLabel", text=f"  {num}. {nome_passo}")

    def _clique_aba_lateral(self, passo_num):
        modo = self.opcao_modo.get()
        if modo == "3" and passo_num in [2, 3]:
            return
        if modo == "2" and passo_num == 3:
            return
        if passo_num <= self.max_passo_alcancado:
            self._exibir_passo(passo_num)

    def _avancar_passo(self, passo_origem):
        modo = self.opcao_modo.get()
        if passo_origem == 2 and modo != "3" and not self.caminho_word:
            messagebox.showwarning("Atenção", "Por favor, selecione um arquivo Word antes de avançar.")
            return

        if passo_origem == 4 and not self.caminhos_pdfs:
            messagebox.showwarning(
                "Atenção", "Por favor, selecione pelo menos um arquivo PDF antes de avançar."
            )
            return

        self._exibir_passo(passo_origem + 1)

    def _voltar_passo(self, passo_origem):
        self._exibir_passo(passo_origem - 1)

    def _selecionar_word(self):
        modo = self.opcao_modo.get()
        if modo == "1":
            caminho = filedialog.asksaveasfilename(
                title="Onde deseja salvar o novo arquivo Word?",
                defaultextension=".docx",
                initialfile="Lista_de_Turmas.docx",
                filetypes=[("Documento Word", "*.docx")],
            )
        else:
            caminho = filedialog.askopenfilename(
                title="Selecione o arquivo Word existente",
                filetypes=[("Documento Word", "*.docx")],
            )

        if caminho:
            self.caminho_word = caminho
            self.lbl_word_selected.config(text=caminho)
            if modo == "2":
                try:
                    dados_temp, _, fonte_det = ler_dados_word_existente(caminho)
                    estilo = detectar_estilo_case(dados_temp)
                    self.opcao_case.set(estilo)
                    self.fonte_detectada_word = fonte_det
                    self.fonte_nome_11.set(fonte_det == 11)
                except Exception:
                    pass
            self._atualizar_resumo()

    def _selecionar_pdfs(self):
        caminhos = filedialog.askopenfilenames(
            title="Selecione um ou mais arquivos PDF",
            filetypes=[("Arquivos PDF", "*.pdf")],
        )

        if caminhos:
            for path in caminhos:
                if path not in self.caminhos_pdfs:
                    self.caminhos_pdfs.append(path)

            self._atualizar_listbox_pdfs()
            self._atualizar_resumo()

    def _remover_pdf_selecionado(self):
        selecao = self.listbox_pdfs.curselection()
        if not selecao:
            messagebox.showwarning("Atenção", "Selecione um PDF na lista para remover.")
            return

        index = selecao[0]
        del self.caminhos_pdfs[index]
        self._atualizar_listbox_pdfs()
        self._atualizar_resumo()

    def _atualizar_listbox_pdfs(self):
        self.listbox_pdfs.delete(0, tk.END)
        for path in self.caminhos_pdfs:
            self.listbox_pdfs.insert(tk.END, os.path.basename(path))

    def _atualizar_resumo(self):
        modo = self.opcao_modo.get()
        if modo == "1":
            self.sum_val_modo.config(text="Criar Novo Documento")
        elif modo == "2":
            self.sum_val_modo.config(text="Atualizar Existente")
        else:
            self.sum_val_modo.config(text="Listar Totais de Turmas")

        if modo == "3":
            self.sum_val_word.config(text="Não necessário")
            self.sum_val_case.config(text="Não aplicável")
            self.sum_val_layout.config(text="Não necessário")
        else:
            if self.caminho_word:
                self.sum_val_word.config(text=os.path.basename(self.caminho_word))
            else:
                self.sum_val_word.config(text="Não selecionado")

            txt_case = "TUDO MAIÚSCULO (AA)" if self.opcao_case.get() == "upper" else "Primeira Maiúscula (Aa)"
            
            if modo == "2":
                tamanho_fonte_txt = f"{self.fonte_detectada_word}pt"
                txt_case = f"Detectado: {txt_case} ({tamanho_fonte_txt})"
            else:
                if self.fonte_nome_11.get():
                    txt_case += " Fonte 11pt"
                else:
                    txt_case += ""

            self.sum_val_case.config(text=txt_case)

            if modo == "2":
                self.sum_val_layout.config(text="Extraído do Word existente")
            else:
                qtd_extras = len(self.colunas_extras)
                if qtd_extras == 0:
                    self.sum_val_layout.config(text="2 Colunas (N.º, Nome)")
                else:
                    cols_txt = ", ".join(self.colunas_extras)
                    self.sum_val_layout.config(text=f"{2 + qtd_extras} Colunas (N.º, Nome, {cols_txt})")

        total_pdfs = len(self.caminhos_pdfs)
        if total_pdfs == 0:
            self.sum_val_pdfs.config(text="Nenhum PDF selecionado")
        else:
            self.sum_val_pdfs.config(text=f"{total_pdfs} arquivo(s) selecionado(s)")

        self._atualizar_sidebar_visual()

    def _log(self, mensagem):
        self.txt_log.config(state="normal")
        self.txt_log.insert(tk.END, mensagem + "\n")
        self.txt_log.see(tk.END)
        self.txt_log.config(state="disabled")

    def _set_progresso(self, valor):
        v_int = int(round(valor))
        self.root.after(0, lambda v=v_int: self._atualizar_interface_progresso(v))

    def _atualizar_interface_progresso(self, v):
        self.progress.configure(value=v)
        self.lbl_porcentagem.config(text=f"{v}%")

    def _executar_processamento(self):
        modo = self.opcao_modo.get()
        if modo != "3" and not self.caminho_word:
            messagebox.showwarning("Erro", "Arquivo Word não selecionado!")
            return
        if not self.caminhos_pdfs:
            messagebox.showwarning("Erro", "Nenhum arquivo PDF selecionado!")
            return

        self.btn_gerar.config(state="disabled")
        self._set_progresso(0)

        threading.Thread(target=self._worker_processamento, daemon=True).start()

    def _worker_processamento(self):
        try:
            dados_consolidados = {}
            modo = self.opcao_modo.get()

            if modo == "3":
                self._log("--- Lendo arquivos PDF para cálculo de totais ---")
                total_pdfs = len(self.caminhos_pdfs)
                for i_pdf, caminho_pdf in enumerate(self.caminhos_pdfs):

                    def pdf_prog_cb(fatia_pag, index=i_pdf):
                        prog_pdf = (index + fatia_pag) / total_pdfs
                        val = prog_pdf * 90
                        self._set_progresso(val)

                    novos_dados = extrair_dados_pdf(
                        caminho_pdf,
                        log_callback=self._log,
                        progress_callback=pdf_prog_cb,
                    )
                    for turma, alunos in novos_dados.items():
                        if turma not in dados_consolidados:
                            dados_consolidados[turma] = {}
                        for num, info in alunos.items():
                            dados_consolidados[turma][num] = info

                self._set_progresso(95)
                self._log("\nCalculando totais dos alunos matriculados...")

                dados_totais = calcular_totais_dados(dados_consolidados)
                self._log("Concluído! Exibindo relatório na tela de totais.")
                self._set_progresso(100)

                self.root.after(
                    0,
                    lambda: self._exibir_janela_relatorio(dados_totais),
                )

            else:
                inicio_pdf = 10 if modo == "2" else 0
                fim_pdf = 70
                font_size_final = 11 if self.fonte_nome_11.get() else 12

                if modo == "2":
                    self._log("Lendo dados do arquivo Word existente...")
                    self._set_progresso(5)
                    dados_consolidados, self.colunas_extras, font_size_detectada = ler_dados_word_existente(self.caminho_word)

                    estilo_word = detectar_estilo_case(dados_consolidados)
                    self.opcao_case.set(estilo_word)
                    self.fonte_detectada_word = font_size_detectada
                    font_size_final = font_size_detectada

                    self.root.after(0, self._atualizar_resumo)
                    self._set_progresso(10)

                total_pdfs = len(self.caminhos_pdfs)
                for i_pdf, caminho_pdf in enumerate(self.caminhos_pdfs):

                    def pdf_prog_cb(fatia_pag, index=i_pdf):
                        prog_pdf = (index + fatia_pag) / total_pdfs
                        val = inicio_pdf + prog_pdf * (fim_pdf - inicio_pdf)
                        self._set_progresso(val)

                    novos_dados = extrair_dados_pdf(
                        caminho_pdf,
                        log_callback=self._log,
                        progress_callback=pdf_prog_cb,
                    )
                    for turma, alunos in novos_dados.items():
                        if turma not in dados_consolidados:
                            dados_consolidados[turma] = {}
                        for num, info in alunos.items():
                            if modo == "2" and num in dados_consolidados[turma]:
                                nome_existente = dados_consolidados[turma][num].get("nome", "")
                                if nome_existente:
                                    info["nome"] = nome_existente
                            dados_consolidados[turma][num] = info

                self._set_progresso(70)

                def word_prog_cb(fatia_turma):
                    val = 70 + fatia_turma * 30
                    self._set_progresso(val)

                gerar_documento_word(
                    dados_consolidados,
                    self.caminho_word,
                    colunas_extras=self.colunas_extras,
                    opcao_case=self.opcao_case.get(),
                    font_size_nome=font_size_final,
                    log_callback=self._log,
                    progress_callback=word_prog_cb,
                )

                self._set_progresso(100)

                self.root.after(
                    0,
                    lambda: messagebox.showinfo(
                        "Sucesso",
                        f"Processo concluído com sucesso!\nSalvo em: {self.caminho_word}",
                    ),
                )
        except Exception as e:
            self.root.after(
                0,
                lambda: messagebox.showerror(
                    "Erro", f"Ocorreu um erro no processamento:\n{str(e)}"
                ),
            )
        finally:
            self.root.after(0, self._finalizar_execucao)

    def _exibir_janela_relatorio(self, dados_totais):
        """Abre uma janela secundária exibindo os totais organizados por seções e agrupados por turno em duas colunas (Manhã e Tarde)."""
        top = tk.Toplevel(self.root)
        top.title("Relatório de Totais das Turmas")
        top.geometry("720x620")
        top.transient(self.root)
        top.grab_set()

        top_bar = ttk.Frame(top, padding="10")
        top_bar.pack(side=tk.TOP, fill=tk.X)

        btn_imprimir = ttk.Button(
            top_bar,
            text="🖨️Imprimir",
            command=lambda: self._imprimir_relatorio_a4(dados_totais),
        )
        btn_imprimir.pack(side=tk.LEFT)

        ttk.Separator(top, orient="horizontal").pack(fill=tk.X, padx=10, pady=(0, 5))

        canvas = tk.Canvas(top, borderwidth=0, highlightthickness=0)
        scrollbar = ttk.Scrollbar(top, orient="vertical", command=canvas.yview)
        scroll_frame = ttk.Frame(canvas, padding="15")

        scroll_frame.bind(
            "<Configure>",
            lambda e: canvas.configure(scrollregion=canvas.bbox("all")),
        )

        canvas_window = canvas.create_window((0, 0), window=scroll_frame, anchor="nw")

        def _on_canvas_configure(event):
            canvas.itemconfig(canvas_window, width=event.width)

        canvas.bind("<Configure>", _on_canvas_configure)
        canvas.configure(yscrollcommand=scrollbar.set)

        def _on_mousewheel(event):
            canvas.yview_scroll(int(-1 * (event.delta / 120)), "units")

        canvas.bind_all("<MouseWheel>", _on_mousewheel)
        top.bind("<Destroy>", lambda e: top.unbind_all("<MouseWheel>"))

        lbl_titulo = ttk.Label(
            scroll_frame,
            text="Relatório de Totais de Alunos Matriculados",
            font=("TkDefaultFont", 12, "bold"),
        )
        lbl_titulo.pack(anchor=tk.W, pady=(0, 10))

        card_geral = ttk.LabelFrame(scroll_frame, text=" Total Geral ", padding="12")
        card_geral.pack(fill=tk.X, pady=(0, 10))

        lbl_total = ttk.Label(
            card_geral,
            text=f"Total Geral de Matriculados: {dados_totais['total_geral']} aluno(s)",
            font=("TkDefaultFont", 11, "bold"),
            foreground="#0056b3",
        )
        lbl_total.pack(anchor=tk.W)

        def criar_secao_cards(parent, titulo, dicionario):
            card = ttk.LabelFrame(parent, text=f" {titulo} ", padding="12")
            card.pack(fill=tk.X, pady=(0, 10))

            if not dicionario:
                ttk.Label(
                    card,
                    text="Nenhum aluno matriculado encontrado.",
                    font=("TkDefaultFont", 9, "italic"),
                ).pack(anchor=tk.W)
                return

            row = 0
            for chave in sorted(dicionario.keys(), key=chave_ordenacao_natural):
                qtd = dicionario[chave]
                lbl_item = ttk.Label(
                    card, text=f"• {chave}:", font=("TkDefaultFont", 9)
                )
                lbl_item.grid(row=row, column=0, sticky=tk.W, padx=(0, 10), pady=2)

                lbl_val = ttk.Label(
                    card, text=f"{qtd} aluno(s)", font=("TkDefaultFont", 9, "bold")
                )
                lbl_val.grid(row=row, column=1, sticky=tk.W, pady=2)
                row += 1

        criar_secao_cards(scroll_frame, "Total por Curso", dados_totais["por_curso"])
        criar_secao_cards(scroll_frame, "Total por Turno", dados_totais["por_turno"])
        criar_secao_cards(scroll_frame, "Total por Série", dados_totais["por_serie"])

        card_turma = ttk.LabelFrame(scroll_frame, text=" Total por Turma ", padding="12")
        card_turma.pack(fill=tk.X, pady=(0, 10))

        por_turma_agrupado = dados_totais.get("por_turma_agrupado", {})

        if not por_turma_agrupado:
            ttk.Label(
                card_turma,
                text="Nenhum aluno matriculado encontrado.",
                font=("TkDefaultFont", 9, "italic"),
            ).pack(anchor=tk.W)
        else:
            dict_manha = {}
            dict_tarde = {}
            dict_outros = {}

            for turno, turmas_dict in por_turma_agrupado.items():
                t_upper = str(turno).upper()
                if "MANH" in t_upper or "MATUTIN" in t_upper:
                    dict_manha[turno] = turmas_dict
                elif "TARD" in t_upper or "VESPERTIN" in t_upper:
                    dict_tarde[turno] = turmas_dict
                else:
                    dict_outros[turno] = turmas_dict

            container_colunas = ttk.Frame(card_turma)
            container_colunas.pack(fill=tk.X, expand=True)

            col_esquerda = ttk.Frame(container_colunas)
            col_esquerda.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 10))

            col_direita = ttk.Frame(container_colunas)
            col_direita.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(10, 0))

            def renderizar_bloco_turno(parent_frame, turno, turmas_dict):
                lbl_header = ttk.Label(
                    parent_frame,
                    text=f"Turno: {turno}",
                    font=("TkDefaultFont", 10, "bold"),
                    foreground="#2F4F4F",
                )
                lbl_header.pack(anchor=tk.W, pady=(4, 2))

                grid_frame = ttk.Frame(parent_frame)
                grid_frame.pack(fill=tk.X, anchor=tk.W, pady=(0, 6))

                for r_idx, turma_nome in enumerate(sorted(turmas_dict.keys(), key=chave_ordenacao_natural)):
                    qtd = turmas_dict[turma_nome]
                    lbl_item = ttk.Label(
                        grid_frame, text=f"• {turma_nome}:", font=("TkDefaultFont", 9)
                    )
                    lbl_item.grid(row=r_idx, column=0, sticky=tk.W, padx=(5, 10), pady=1)

                    lbl_val = ttk.Label(
                        grid_frame, text=f"{qtd} aluno(s)", font=("TkDefaultFont", 9, "bold")
                    )
                    lbl_val.grid(row=r_idx, column=1, sticky=tk.W, pady=1)

            if dict_manha:
                for turno in sorted(dict_manha.keys(), key=ordenar_turnos):
                    renderizar_bloco_turno(col_esquerda, turno, dict_manha[turno])
            else:
                lbl_vazio_m = ttk.Label(
                    col_esquerda,
                    text="Turno: Manhã\n(Nenhuma turma)",
                    font=("TkDefaultFont", 9, "italic"),
                    foreground="#777777",
                )
                lbl_vazio_m.pack(anchor=tk.W)

            if dict_tarde:
                for turno in sorted(dict_tarde.keys(), key=ordenar_turnos):
                    renderizar_bloco_turno(col_direita, turno, dict_tarde[turno])
            else:
                lbl_vazio_t = ttk.Label(
                    col_direita,
                    text="Turno: Tarde\n(Nenhuma turma)",
                    font=("TkDefaultFont", 9, "italic"),
                    foreground="#777777",
                )
                lbl_vazio_t.pack(anchor=tk.W)

            if dict_outros:
                ttk.Separator(card_turma, orient="horizontal").pack(fill=tk.X, pady=8)
                frame_outros = ttk.Frame(card_turma)
                frame_outros.pack(fill=tk.X)

                for turno in sorted(dict_outros.keys(), key=ordenar_turnos):
                    renderizar_bloco_turno(frame_outros, turno, dict_outros[turno])

        canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def _imprimir_relatorio_a4(self, dados_totais):
        """Gera um relatório HTML formatado em 12pt com turmas agrupadas por turno e solicita ao usuário escolher a impressora."""
        try:
            data_hora_geracao = datetime.now().strftime("%d/%m/%Y às %H:%M:%S")

            def montar_tabela_html(dicionario):
                if not dicionario:
                    return "<p class='vazio'>Nenhum aluno matriculado encontrado.</p>"
                rows = []
                for chave in sorted(dicionario.keys(), key=chave_ordenacao_natural):
                    qtd = dicionario[chave]
                    rows.append(
                        f"<tr><td class='label'>• {chave}</td><td class='val'>{qtd}</td></tr>"
                    )
                return f"<table>{''.join(rows)}</table>"

            def montar_turmas_agrupadas_html(por_turma_agrupado):
                if not por_turma_agrupado:
                    return "<p class='vazio'>Nenhum aluno matriculado encontrado.</p>"

                blocks = []
                turnos_ordenados = sorted(por_turma_agrupado.keys(), key=ordenar_turnos)

                for turno in turnos_ordenados:
                    turmas_dict = por_turma_agrupado[turno]
                    rows = []
                    for turma_nome in sorted(turmas_dict.keys(), key=chave_ordenacao_natural):
                        qtd = turmas_dict[turma_nome]
                        rows.append(
                            f"<tr><td class='label'>• {turma_nome}</td><td class='val'>{qtd}</td></tr>"
                        )

                    group_html = f"""
                    <div class="turno-block">
                        <div class="turno-header">Turno: {turno}</div>
                        <table>{''.join(rows)}</table>
                    </div>
                    """
                    blocks.append(group_html)

                return "".join(blocks)

            html_content = f"""<!DOCTYPE html>
<html lang="pt-BR">
<head>
    <meta charset="UTF-8">
    <title>Relatório de Totais de Turmas</title>
    <style>
        @page {{
            size: A4 portrait;
            margin: 1.5cm;
        }}
        body {{
            font-family: Arial, Helvetica, sans-serif;
            color: black;
            margin: 0;
            padding: 0;
            font-size: 12pt;
            background: #fff;
        }}
        .header {{
            text-align: center;
            border-bottom: 2px solid #2f4f4f;
            padding-bottom: 6px;
            margin-bottom: 12px;
        }}
        .header h1 {{
            margin: 0;
            font-size: 16pt;
            color: #2f4f4f;
            text-transform: uppercase;
        }}
        .header p {{
            margin: 4px 0 0 0;
            font-size: 11pt;
            color: #555;
            font-weight: bold;
        }}
        .total-banner {{
            background: #f0f4f8;
            border: 1px solid #b8cce4;
            border-radius: 4px;
            padding: 8px 12px;
            margin-bottom: 12px;
            text-align: center;
            font-size: 13pt;
            font-weight: bold;
            color: #0056b3;
        }}
        .columns-container {{
            display: flex;
            gap: 14px;
            align-items: flex-start;
        }}
        .col-left, .col-right {{
            flex: 1;
            display: flex;
            flex-direction: column;
            gap: 12px;
        }}
        .card {{
            border: 1px solid #ccc;
            border-radius: 4px;
            padding: 10px 12px;
            background: #fff;
            page-break-inside: avoid;
            break-inside: avoid;
        }}
        .card-title {{
            font-weight: bold;
            font-size: 13pt;
            color: #2f4f4f;
            border-bottom: 1.5px solid #2f4f4f;
            padding-bottom: 4px;
            margin-bottom: 8px;
            text-transform: uppercase;
        }}
        .turno-block {{
            margin-bottom: 10px;
        }}
        .turno-header {{
            font-weight: bold;
            font-size: 12pt;
            color: #0056b3;
            background-color: #eef5fc;
            padding: 3px 6px;
            border-left: 4px solid #0056b3;
            margin-top: 6px;
            margin-bottom: 4px;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
        }}
        td {{
            padding: 3px 6px;
            vertical-align: middle;
            border-bottom: 1px solid #eee;
            font-size: 12pt;
        }}
        td.label {{
            font-weight: normal;
        }}
        td.val {{
            font-weight: bold;
            text-align: right;
            white-space: nowrap;
        }}
        p.vazio {{
            font-style: italic;
            color: #777;
            margin: 0;
            font-size: 11pt;
        }}
    </style>
</head>
<body onload="window.print();">
    <div class="header">
        <h1>Relatório de Totais de Alunos Matriculados</h1>
        <p>Gerado em: {data_hora_geracao}</p>
    </div>

    <div class="total-banner">
        Total Geral de Matriculados: {dados_totais['total_geral']} aluno(s)
    </div>

    <div class="columns-container">
        <div class="col-left">
            <div class="card">
                <div class="card-title">Total por Curso</div>
                {montar_tabela_html(dados_totais['por_curso'])}
            </div>

            <div class="card">
                <div class="card-title">Total por Turno</div>
                {montar_tabela_html(dados_totais['por_turno'])}
            </div>

            <div class="card">
                <div class="card-title">Total por Série</div>
                {montar_tabela_html(dados_totais['por_serie'])}
            </div>
        </div>

        <div class="col-right">
            <div class="card">
                <div class="card-title">Total por Turma</div>
                {montar_turmas_agrupadas_html(dados_totais.get('por_turma_agrupado', {}))}
            </div>
        </div>
    </div>
</body>
</html>"""

            temp_dir = tempfile.gettempdir()
            caminho_html = os.path.join(temp_dir, "relatorio_totais_a4.html")

            if caminho_html not in self.arquivos_temporarios:
                self.arquivos_temporarios.append(caminho_html)

            with open(caminho_html, "w", encoding="utf-8") as f:
                f.write(html_content)

            messagebox.showinfo(
                "Seleção de Impressora",
                "O relatório será exibido no seu navegador para que você possa escolher em qual impressora imprimir.",
            )
            webbrowser.open(f"file://{os.path.abspath(caminho_html)}")

        except Exception as e:
            messagebox.showerror(
                "Erro ao Imprimir",
                f"Não foi possível processar o relatório de impressão:\n{str(e)}",
            )

    def _finalizar_execucao(self):
        self.btn_gerar.config(state="normal")

    def reiniciar(self):
        self.opcao_modo.set("1")
        self.opcao_case.set("upper")
        self.fonte_nome_11.set(False)
        self.fonte_detectada_word = 12
        self.caminho_word = ""
        self.colunas_extras = []
        self.caminhos_pdfs = []
        self.entry_nova_coluna.delete(0, tk.END)

        self.lbl_word_selected.config(text="Nenhum arquivo selecionado.")
        self._atualizar_botoes_case()
        self._atualizar_listbox_colunas()
        self._atualizar_listbox_pdfs()

        self.passo_atual = 1
        self.max_passo_alcancado = 1

        self._set_progresso(0)
        self.txt_log.config(state="normal")
        self.txt_log.delete("1.0", tk.END)
        self.txt_log.config(state="disabled")

        self._exibir_passo(1)
        self._atualizar_resumo()

if __name__ == "__main__":
    root = tk.Tk()
    app = SideTabWizardApp(root)
    root.mainloop()